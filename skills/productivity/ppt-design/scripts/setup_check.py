#!/usr/bin/env python3
"""Environment check and one-click dependency install for luobo-ppt.

Usage:
  python setup_check.py                 check only, print a report
  python setup_check.py --install       check, then install what is missing
  python setup_check.py --quick         exit at once if a previous check passed for this version
  python setup_check.py --selftest      also build a tiny deck end to end
  python setup_check.py --json          machine-readable report

Typical calls
  installer (install.sh / install.ps1):   setup_check.py --install --selftest
  first use inside an agent (SKILL.md):   setup_check.py --quick --install

What is checked
  required     Python >= 3.8, python-pptx, Pillow, numpy, lxml, a CJK font for text measuring
  recommended  LibreOffice (previews + visual QA), pdftoppm or PyMuPDF (PDF -> PNG),
               scipy (background removal for AI images), requests (AI image APIs)
  optional     rembg (better cut-outs for photographic AI images), image model config
The result is cached in ~/.config/luobo-ppt/setup.json.
Exit code 0 when all required items are present.
"""
import argparse
import importlib
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

try:  # make Chinese output safe on Windows consoles
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

PY_REQUIRED = [("pptx", "python-pptx"), ("PIL", "Pillow"), ("numpy", "numpy"), ("lxml", "lxml")]
PY_RECOMMENDED = [("scipy", "scipy"), ("requests", "requests"), ("pymupdf", "pymupdf")]
PY_OPTIONAL = [("rembg", "rembg")]

OS = platform.system()  # Linux / Darwin / Windows


def say(msg=""):
    print(msg, flush=True)


# ---------------------------------------------------------------- checks
def has_module(mod):
    try:
        importlib.invalidate_caches()
        importlib.import_module(mod)
        return True
    except Exception:
        return False


def check_all():
    import common
    importlib.reload(common)
    r = {"version": common.SKILL_VERSION, "python": sys.executable,
         "python_version": platform.python_version(), "os": f"{OS} {platform.release()}",
         "items": []}

    def item(name, level, ok, detail="", fix=""):
        r["items"].append({"name": name, "level": level, "ok": bool(ok), "detail": detail, "fix": fix})

    item("Python >= 3.8", "required", sys.version_info >= (3, 8), platform.python_version(),
         "install Python 3.8 or newer")
    for mod, pkg in PY_REQUIRED:
        item(pkg, "required", has_module(mod), "", f"pip install {pkg}")
    font = common.find_cjk_font()
    item("CJK font (for text measuring)", "required", bool(font), font or "",
         {"Linux": "apt-get install fonts-noto-cjk", "Darwin": "bundled with the system, usually nothing to do",
          "Windows": "微软雅黑 is bundled with the system, usually nothing to do"}.get(OS, "install Noto Sans CJK"))
    soffice = common.find_soffice()
    item("LibreOffice", "recommended", bool(soffice), soffice or "",
         {"Linux": "apt-get install libreoffice-impress", "Darwin": "brew install --cask libreoffice",
          "Windows": "winget install TheDocumentFoundation.LibreOffice"}.get(OS, "install LibreOffice"))
    pdftoppm = common.find_pdftoppm()
    fitz_ok = has_module("pymupdf") or has_module("fitz")
    item("PDF to image (pdftoppm or PyMuPDF)", "recommended", bool(pdftoppm) or fitz_ok,
         pdftoppm or ("PyMuPDF" if fitz_ok else ""), "pip install pymupdf")
    for mod, pkg in PY_RECOMMENDED:
        if pkg == "pymupdf":
            continue
        item(pkg, "recommended", has_module(mod), "", f"pip install {pkg}")
    for mod, pkg in PY_OPTIONAL:
        item(pkg, "optional", has_module(mod), "", f"pip install {pkg} (optional)")
    try:
        from imagegen import find_config, ImageGen
        cfg = find_config()
        ok, why = ImageGen(cfg).available()
        item("Image model config", "optional", ok, why if ok else f"not configured, will use procedural illustrations ({why})",
             "see references/imagegen_config.md")
    except Exception as e:  # noqa: BLE001
        item("Image model config", "optional", False, f"not checked: {e}", "")
    r["soffice"] = soffice
    r["pdftoppm"] = pdftoppm
    r["cjk_font"] = font
    r["ready"] = all(i["ok"] for i in r["items"] if i["level"] == "required")
    r["full"] = r["ready"] and all(i["ok"] for i in r["items"] if i["level"] == "recommended")
    return r


# ---------------------------------------------------------------- installers
def run(cmd, timeout=1800, check=False):
    say("  $ " + " ".join(cmd))
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    except Exception as e:  # noqa: BLE001
        say(f"    failed: {e}")
        return False, str(e)
    out = (p.stdout or "") + (p.stderr or "")
    if p.returncode != 0:
        tail = out.strip().splitlines()[-3:]
        say("    failed: " + " | ".join(tail))
    return p.returncode == 0, out


def pip_install(pkgs):
    if not pkgs:
        return True
    ok, _ = run([sys.executable, "-m", "pip", "--version"])
    if not ok:
        ok, _ = run([sys.executable, "-m", "ensurepip", "--upgrade"])
        if not ok and OS == "Linux" and shutil.which("apt-get") and can_sudo():
            pre = [] if os.geteuid() == 0 else ["sudo", "-n"]
            run(pre + ["env", "DEBIAN_FRONTEND=noninteractive", "apt-get", "install", "-y", "-qq", "python3-pip"])
    base = [sys.executable, "-m", "pip", "install", "--disable-pip-version-check", "-q"]
    ok, out = run(base + pkgs)
    if ok:
        return True
    if "externally-managed" in out or "EXTERNALLY-MANAGED" in out:
        ok, out = run(base + ["--break-system-packages"] + pkgs)
        if ok:
            return True
    if "Permission" in out or "permission" in out or "EnvironmentError" in out:
        ok, out = run(base + ["--user"] + pkgs)
        if ok:
            return True
        ok, out = run(base + ["--user", "--break-system-packages"] + pkgs)
    return ok


def can_sudo():
    if OS != "Linux":
        return False
    if hasattr(os, "geteuid") and os.geteuid() == 0:
        return True
    if shutil.which("sudo"):
        ok, _ = run(["sudo", "-n", "true"], timeout=10)
        return ok
    return False


def sys_install(want):
    """want: set of 'libreoffice', 'poppler', 'cjkfont'. Best effort, never raises."""
    if not want:
        return
    if OS == "Linux":
        pre = [] if (hasattr(os, "geteuid") and os.geteuid() == 0) else ["sudo", "-n"]
        if not can_sudo():
            say("  no root/sudo access, skipping system component install; run the suggested commands above manually")
            return
        if shutil.which("apt-get"):
            pk = {"libreoffice": ["libreoffice-impress"], "poppler": ["poppler-utils"], "cjkfont": ["fonts-noto-cjk"]}
            pkgs = sum((pk[w] for w in want), [])
            env_pre = pre + ["env", "DEBIAN_FRONTEND=noninteractive"]
            run(env_pre + ["apt-get", "update", "-qq"], timeout=600)
            run(env_pre + ["apt-get", "install", "-y", "-qq", "--no-install-recommends"] + pkgs)
        elif shutil.which("dnf") or shutil.which("yum"):
            mgr = "dnf" if shutil.which("dnf") else "yum"
            pk = {"libreoffice": ["libreoffice-impress"], "poppler": ["poppler-utils"],
                  "cjkfont": ["google-noto-sans-cjk-fonts"]}
            run(pre + [mgr, "install", "-y"] + sum((pk[w] for w in want), []))
        elif shutil.which("pacman"):
            pk = {"libreoffice": ["libreoffice-fresh"], "poppler": ["poppler"], "cjkfont": ["noto-fonts-cjk"]}
            run(pre + ["pacman", "-S", "--noconfirm"] + sum((pk[w] for w in want), []))
        else:
            say("  unrecognized package manager, please install manually")
    elif OS == "Darwin":
        if not shutil.which("brew"):
            say("  Homebrew not found. LibreOffice can be downloaded from https://www.libreoffice.org/download/")
            return
        if "libreoffice" in want:
            run(["brew", "install", "--cask", "libreoffice"])
        if "poppler" in want:
            run(["brew", "install", "poppler"])
    elif OS == "Windows":
        if "libreoffice" in want:
            if shutil.which("winget"):
                run(["winget", "install", "-e", "--id", "TheDocumentFoundation.LibreOffice",
                     "--accept-source-agreements", "--accept-package-agreements", "--silent"])
            elif shutil.which("choco"):
                run(["choco", "install", "libreoffice-fresh", "-y"])
            else:
                say("  winget/choco not found. LibreOffice can be downloaded from https://www.libreoffice.org/download/")


def install_missing(report):
    miss = {i["name"] for i in report["items"] if not i["ok"]}
    pkgs = [pkg for mod, pkg in PY_REQUIRED + PY_RECOMMENDED if pkg in miss]
    if "PDF to image (pdftoppm or PyMuPDF)" in miss and "pymupdf" not in pkgs:
        pkgs.append("pymupdf")
    if pkgs:
        say(f"\nInstalling Python packages: {' '.join(pkgs)}")
        pip_install(pkgs)
    want = set()
    if "LibreOffice" in miss:
        want.add("libreoffice")
    if "CJK font (for text measuring)" in miss:
        want.add("cjkfont")
    if want:
        say(f"\nInstalling system components: {', '.join(sorted(want))} (may take a few minutes)")
        sys_install(want)


# ---------------------------------------------------------------- selftest
def selftest(report):
    say("\nSelf-test: build a two-page sample with procedural illustrations ...")
    tmp = Path(tempfile.mkdtemp(prefix="pb_selftest_"))
    try:
        from pptx import Presentation
        p = Presentation()
        s = p.slides.add_slide(p.slide_layouts[0])
        s.shapes.title.text = "Self-test cover"
        s.placeholders[1].text = "luobo-ppt"
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = "Three key points"
        tf = s.placeholders[1].text_frame
        tf.text = "First point"
        for t in ("Second point", "Third point"):
            tf.add_paragraph().text = t
        src = tmp / "src.pptx"
        p.save(str(src))
        W = tmp / "w"
        py = sys.executable
        steps = [[py, HERE / "extract_deck.py", src, W, "--no-render"]]
        for st in steps:
            if subprocess.run([str(x) for x in st], capture_output=True).returncode:
                raise RuntimeError("extract_deck failed")
        deck = {"meta": {"title": "Test", "series": "", "anchor": {"description": "a glossy app icon", "glyph": "T"}},
                "slides": [{"n": 1, "type": "cover", "title": "Self-test cover", "subtitle": "luobo-ppt"},
                           {"n": 2, "type": "list", "title": "Three key points",
                            "items": [{"text": "First point"}, {"text": "Second point"}, {"text": "Third point"}]}]}
        (W / "deck.json").write_text(json.dumps(deck, ensure_ascii=False), encoding="utf-8")
        (W / "style.json").write_text(json.dumps({"name": "selftest"}), encoding="utf-8")
        for st in ([py, HERE / "plan_deck.py", W, "--level", "standard"],
                   [py, HERE / "make_assets.py", W, "--procedural"],
                   [py, HERE / "build_deck.py", W],
                   [py, HERE / "qa_check.py", W]):
            pr = subprocess.run([str(x) for x in st], capture_output=True, text=True)
            if pr.returncode:
                raise RuntimeError(f"{Path(str(st[1])).name} failed: {(pr.stdout + pr.stderr).strip()[-300:]}")
        out = W / "out" / "beautified.pptx"
        Presentation(str(out))
        say("  assembly and automated checks passed")
        if report.get("soffice"):
            pr = subprocess.run([py, str(HERE / "render.py"), str(out), str(W / "out" / "png")],
                                capture_output=True, text=True, timeout=300)
            say("  render preview " + ("passed" if pr.returncode == 0 else "failed: " + (pr.stdout + pr.stderr).strip()[-200:]))
        return True
    except Exception as e:  # noqa: BLE001
        say(f"  self-test failed: {e}")
        return False
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ---------------------------------------------------------------- report
def print_report(r):
    mark = {True: "OK ", False: "-- "}
    say(f"\nluobo-ppt {r['version']} environment check   Python {r['python_version']}   {r['os']}")
    for level, title in (("required", "required"), ("recommended", "recommended"), ("optional", "optional")):
        say(f"\n[{title}]")
        for i in r["items"]:
            if i["level"] != level:
                continue
            line = f"  {mark[i['ok']]} {i['name']}"
            if i["detail"]:
                line += f"   {i['detail']}"
            if not i["ok"] and i["fix"] and level != "optional":
                line += f"\n       suggestion: {i['fix']}"
            say(line)
    say("")
    if r["full"]:
        say("Verdict: everything is ready.")
    elif r["ready"]:
        say("Verdict: usable. When recommended components are missing, the beautified PPT is still built, just without preview images or AI cut-outs.")
    else:
        say("Verdict: required components are missing, install them first (add --install to install automatically).")


def save_cache(r):
    import common
    try:
        common.SETUP_CACHE.parent.mkdir(parents=True, exist_ok=True)
        data = dict(r, checked_at=time.strftime("%Y-%m-%d %H:%M:%S"))
        common.SETUP_CACHE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    except Exception:
        pass


def main():
    ap = argparse.ArgumentParser(description="luobo-ppt environment check")
    ap.add_argument("--install", action="store_true", help="install missing dependencies")
    ap.add_argument("--quick", action="store_true", help="skip if a previous check passed")
    ap.add_argument("--selftest", action="store_true", help="build a tiny deck end to end")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()

    if a.quick:
        try:
            import common
            c = json.loads(common.SETUP_CACHE.read_text(encoding="utf-8"))
            if c.get("ready") and c.get("version") == common.SKILL_VERSION and c.get("python") == sys.executable:
                say(f"luobo-ppt environment is ready (check passed at {c.get('checked_at', '')}).")
                return 0
        except Exception:
            pass

    missing_core = [pkg for mod, pkg in PY_REQUIRED if not has_module(mod)]
    if missing_core and a.install:
        say(f"Installing required Python packages: {' '.join(missing_core)}")
        pip_install(missing_core)
    if [pkg for mod, pkg in PY_REQUIRED if not has_module(mod)] and not a.install:
        # cannot even import common helpers reliably; report minimal info
        pass

    r = check_all()
    if a.install and not r["full"]:
        install_missing(r)
        r = check_all()
    if a.selftest and r["ready"]:
        r["selftest"] = selftest(r)
        if not r["selftest"]:
            r["ready"] = False
    if a.json:
        print(json.dumps(r, ensure_ascii=False, indent=2))
    else:
        print_report(r)
    save_cache(r)
    return 0 if r["ready"] else 1


if __name__ == "__main__":
    sys.exit(main())
