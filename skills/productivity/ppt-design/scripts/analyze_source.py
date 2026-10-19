#!/usr/bin/env python3
"""Generation step 1: measure the source material and suggest a page count.

Usage:
  python analyze_source.py WORK_DIR --file article.md [--file notes.docx ...]
  python analyze_source.py WORK_DIR --topic "What are Agent Skills and how do ordinary people use them"
  options: --scenario oral|talk|reading|report   (default oral: oral video)
           --minutes 3                            target speaking time, overrides length-based estimate

Writes WORK_DIR/analysis.json and WORK_DIR/analysis.md, and copies the source text to
WORK_DIR/source.md. Claude reads analysis.md, then writes deck.json
(see references/generate.md).

What it measures: length (CJK characters + latin words), headings, numbers that look like
data points (candidates for stats slides), keyword signals for picking a style preset, and a
page-count range for the scenario.
"""
import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import save_json, load_json, SKILL_DIR, style_file  # noqa: E402

# on-screen pace per scenario: how much source text one slide usually carries
SCENARIOS = {
    "oral":    {"label": "Oral video", "chars_per_slide": 150, "sec_per_slide": 25, "density": "low"},
    "talk":    {"label": "Live talk", "chars_per_slide": 180, "sec_per_slide": 60, "density": "low"},
    "reading": {"label": "Reading", "chars_per_slide": 260, "sec_per_slide": 0, "density": "high"},
    "report":  {"label": "Report", "chars_per_slide": 220, "sec_per_slide": 90, "density": "high"},
}


def read_any(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".md", ".txt", ".markdown", ""):
        return p.read_text(encoding="utf-8", errors="ignore")
    if ext == ".docx":
        try:
            import docx  # python-docx
            d = docx.Document(str(p))
            out = []
            for para in d.paragraphs:
                t = para.text.strip()
                if not t:
                    continue
                style = (para.style.name or "").lower()
                if style.startswith("heading"):
                    lvl = re.findall(r"\d", style)
                    out.append("#" * int(lvl[0] if lvl else 2) + " " + t)
                else:
                    out.append(t)
            return "\n\n".join(out)
        except ImportError:
            raise SystemExit("reading .docx requires python-docx: pip install python-docx")
    if ext == ".pdf":
        try:
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz
            return "\n\n".join(page.get_text() for page in fitz.open(str(p)))
        except ImportError:
            raise SystemExit("reading .pdf requires PyMuPDF: pip install pymupdf")
    if ext == ".pptx":
        from pptx import Presentation
        out = []
        for i, s in enumerate(Presentation(str(p)).slides, 1):
            out.append(f"## Page {i}")
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    out.append(sh.text_frame.text.strip())
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
                out.append("Notes: " + s.notes_slide.notes_text_frame.text.strip())
        return "\n\n".join(out)
    return p.read_text(encoding="utf-8", errors="ignore")


def measure(text):
    cjk = len(re.findall(r"[一-鿿]", text))
    latin_words = len(re.findall(r"[A-Za-z][A-Za-z0-9'\-]*", text))
    units = cjk + latin_words  # rough reading units
    headings = [(len(m.group(1)), m.group(2).strip()) for m in re.finditer(r"^(#{1,4})\s+(.+)$", text, re.M)]
    paras = [p for p in re.split(r"\n\s*\n", text) if p.strip() and not p.strip().startswith("#")]
    # data points: numbers with a unit or %, keep the sentence fragment around them
    data = []
    for m in re.finditer(r"[^，。；\n]{0,18}?\d[\d,\.]*\s*(?:%|万|亿|倍|天|小时|分钟|秒|年|月|次|个|人|元|美元|k|K|M|B)[^，。；\n]{0,12}", text):
        frag = m.group(0).strip()
        if frag and frag not in data:
            data.append(frag)
    lang = "zh" if cjk >= latin_words else "en"
    return {"cjk_chars": cjk, "latin_words": latin_words, "units": units, "headings": headings,
            "paragraphs": len(paras), "data_points": data[:30], "language": lang}


def style_signals(text):
    idx_path = style_file("index")
    if not idx_path.exists():
        return []
    idx = load_json(idx_path)
    scores = []
    for st in idx.get("styles", []):
        hits = [k for k in st.get("keywords", []) if k.lower() in text.lower()]
        if hits:
            scores.append({"style": st["id"], "hits": hits, "score": sum(len(h) for h in hits)})
    return sorted(scores, key=lambda s: -s["score"])


def page_range(units, scenario, minutes=None):
    sc = SCENARIOS[scenario]
    if minutes and sc["sec_per_slide"]:
        mid = max(4, round(minutes * 60 / sc["sec_per_slide"]))
    elif units:
        mid = max(4, round(units / sc["chars_per_slide"]) + 2)  # +2 for cover and closing
    else:
        mid = 8
    lo, hi = max(4, int(mid * 0.8)), max(5, int(round(mid * 1.2)))
    return lo, mid, hi


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("work")
    ap.add_argument("--file", action="append", default=[])
    ap.add_argument("--topic")
    ap.add_argument("--scenario", default="oral", choices=list(SCENARIOS))
    ap.add_argument("--minutes", type=float)
    a = ap.parse_args()
    work = Path(a.work)
    work.mkdir(parents=True, exist_ok=True)
    texts = [read_any(f) for f in a.file]
    text = "\n\n".join(texts)
    if a.topic:
        text = (f"# {a.topic}\n\n" + text) if text else f"# {a.topic}\n"
    (work / "source.md").write_text(text, encoding="utf-8")
    m = measure(text)
    topic_only = not a.file
    lo, mid, hi = page_range(0 if topic_only else m["units"], a.scenario, a.minutes)
    h2 = sum(1 for lv, _ in m["headings"] if lv == 2)
    if h2 >= 3 and not a.minutes:  # an outline: at least one slide per section plus a cover
        mid = max(mid, h2 + 1)
        lo, hi = max(lo, h2), max(hi, h2 + 3)
    sc = SCENARIOS[a.scenario]
    res = {"scenario": a.scenario, "scenario_label": sc["label"], "density": sc["density"],
           "topic_only": topic_only, "pages": {"min": lo, "suggested": mid, "max": hi},
           "style_signals": style_signals(text), **m}
    save_json(res, work / "analysis.json")

    lines = [f"# Source analysis", "",
             f"- Scenario: {sc['label']} (default density {sc['density']})",
             "- Source: " + ("topic only, build the content yourself; verify any facts and figures first" if topic_only
                         else "{} CJK characters, {} English words, {} paragraphs".format(m["cjk_chars"], m["latin_words"], m["paragraphs"])),
             f"- Suggested pages: {lo} to {hi}, recommended {mid} (including cover and closing)",
             f"- Language: {m['language']}"]
    if a.minutes:
        lines.append(f"- Target duration: {a.minutes:g} minutes")
    if m["headings"]:
        lines += ["", "## Source heading structure", ""] + [f"{'  ' * (lv - 1)}- {t}" for lv, t in m["headings"]]
    if m["data_points"]:
        lines += ["", "## Numbers that could become data slides", ""] + [f"- {d}" for d in m["data_points"][:15]]
    if res["style_signals"]:
        lines += ["", "## Style signals", ""] + [f"- {s['style']}: matched {', '.join(s['hits'])}" for s in res["style_signals"][:4]]
    (work / "analysis.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))
    print(f"\nwrote {work/'analysis.json'}, {work/'analysis.md'}, {work/'source.md'}")


if __name__ == "__main__":
    main()
